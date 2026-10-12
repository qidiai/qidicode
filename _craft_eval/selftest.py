#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""selftest.py — T1 自测 (规格书 §6).

不依赖真 CLI、不读真语料内容（只对语料做路径/大小计数）。
用 python-docx 1.2.0 造 4 个 fixture 验证 truth/norm/metrics/sampling。
全过打印 SELFTEST OK 退出码 0，任何断言失败退出码 1。
"""
import io
import os
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from craft_eval import metrics, sampling, truth  # noqa: E402
import run_w2  # noqa: E402  (写路径防护 / _verdict 的回归测试)

FIX_DIR = os.path.join(HERE, "out", "_selftest")


def setup_stdout():
    try:
        if getattr(sys.stdout, "_craft_utf8", False):
            return
        if sys.stdout is not None and hasattr(sys.stdout, "buffer"):
            w = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
            w._craft_utf8 = True
            sys.stdout = w
    except Exception:
        pass


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print("  ok  - %s" % msg)


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
def _make_plain(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("第一段 ABC 123")
    doc.add_paragraph("第二段 中文内容")
    doc.add_paragraph("第三段 tail")
    doc.save(path)
    return path


def _make_table(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("前段落")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "R1C1"
    t.cell(0, 1).text = "R1C2"
    t.cell(1, 0).text = "R2C1"
    t.cell(1, 1).text = "R2C2"
    doc.save(path)
    return path


def _make_fullwidth(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("全角\u3000空格\u3000测试")
    doc.add_paragraph("中文ABC")
    doc.save(path)
    return path


def _inject_del(path, tag):
    """在 word/document.xml 的 <w:body> 末尾注入 <w:del><w:r><w:t>tag</w:t></w:r></w:del>。"""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                s = data.decode("utf-8")
                inject = ('<w:del><w:r><w:t xml:space="preserve">%s</w:t>'
                          '</w:r></w:del>' % tag)
                check_s = s
                idx = check_s.find("</w:body>")
                if idx < 0:
                    raise AssertionError("document.xml 无 </w:body>")
                s = check_s[:idx] + inject + check_s[idx:]
                data = s.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)
    return path


def _make_del_fixture(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("保留文本")
    doc.save(path)
    _inject_del(path, "删除文本XYZ")
    return path


def _para_full_text(path):
    from docx import Document
    return "".join(p.text for p in Document(path).paragraphs)


# --------------------------------------------------------------------------- #
# 断言
# --------------------------------------------------------------------------- #
def test_truth_plain(p1):
    print("[1] truth 纯段落 == python-docx 段落全文")
    t = truth.extract_truth(p1)
    check(t["error"] is None, "extract_truth error is None")
    expect = _para_full_text(p1)
    check(t["raw"] == expect, "raw == python-docx 全文 (%r == %r)" % (t["raw"], expect))
    check(t["norm_chars"] == len(metrics.norm(t["raw"])), "norm_chars 一致")
    check(t["del_chars"] == 0, "纯段落 del_chars == 0")


def test_truth_table(p2):
    print("[2] truth 含表格")
    t = truth.extract_truth(p2)
    check(t["error"] is None, "error is None")
    for cell in ("R1C1", "R1C2", "R2C1", "R2C2", "前段落"):
        check(cell in t["raw"], "真值含表格/段落文本 %r" % cell)


def test_del(p4):
    print("[3] w:del 内文本被剔除")
    t = truth.extract_truth(p4)
    check(t["error"] is None, "error is None")
    check("删除文本XYZ" not in t["raw"], "删除修订文本被剔除出 raw")
    check("保留文本" in t["raw"], "非删除文本仍在 raw")
    check(t["del_chars"] == len("删除文本XYZ"),
          "del_chars == %d (got %d)" % (len("删除文本XYZ"), t["del_chars"]))
    check(t["del_chars"] > 0, "del_chars > 0")


def test_norm(p3):
    print("[4] norm 计数（含全角空格清除）")
    t = truth.extract_truth(p3)
    check(t["error"] is None, "error is None")
    n = metrics.norm(t["raw"])
    check("\u3000" not in n, "全角空格 U+3000 已被清除")
    check(len(n) == t["norm_chars"], "norm_chars == len(norm(raw))")
    check(n == "全角空格测试中文ABC", "norm 结果正确 (%r)" % n)


def test_metrics():
    print("[5] metrics")
    check(metrics.u_fffd_count("abc\uFFFD\uFFFD") == 2, "u_fffd_count 计数正确")
    check(metrics.u_fffd_count("正常中文") == 0, "无替换字符为 0")
    same = metrics.completeness("abc def", 6)  # norm -> "abcdef" len 6
    check(same["ratio"] == 1.0, "completeness 同文本 == 1.0")
    check(same["craft_norm"] == 6, "craft_norm == 6")
    miss = metrics.completeness("abd", 6)
    check(miss["ratio"] < 1.0, "缺字 completeness < 1.0 (%s)" % miss["ratio"])
    check(0.0 <= metrics.cjk_ratio("中文abc") <= 1.0, "cjk_ratio 在 [0,1]")


def test_sampling():
    print("[6] sampling")
    m1 = sampling.build_manifest(sampling.DEFAULT_CORPUS, seed=20261012)
    m2 = sampling.build_manifest(sampling.DEFAULT_CORPUS, seed=20261012)
    check(m1 == m2, "同 seed 两次 build_manifest 完全一致")

    a, _b, _c, d = sampling.pools(sampling.DEFAULT_CORPUS)
    check(all(s <= sampling.BIG_BYTES for _p, s in a), "A 池无 >50MB")
    check(all("业绩" not in p for p, _s in a), "A 池无 '业绩' 路径")

    top5 = [p for p, _s in d[:5]]
    d_entries = [e for e in m1 if e["tier"] == "D"]
    d_manifest = [e["path"] for e in d_entries]
    check(d_manifest == top5, "D 组按大小前 5")
    check(all(d_entries[i]["size_mb"] >= d_entries[i + 1]["size_mb"]
              for i in range(len(d_entries) - 1)), "D 组大小降序")
    return a, _b, _c, d


def test_pool_counts(a, b, c, d):
    print("[7] 池计数（只读路径/大小，不解析文件内容）")
    print("  A pool=%d  B pool=%d  C pool=%d  D pool=%d" % (len(a), len(b), len(c), len(d)))
    print("  (规格书参考 A≈502/B≈839/C≈52/D=25 为未排除>50MB 的原始计数；"
          "按排除规则 A/B/C 略小)")
    check(len(a) >= sampling.A_QUOTA, "A 池 >= 配额 25")
    check(len(b) >= sampling.B_QUOTA, "B 池 >= 配额 10")
    check(len(c) >= sampling.C_QUOTA, "C 池 >= 配额 10")
    check(len(d) >= sampling.D_QUOTA, "D 池 >= 配额 5")


def test_write_guard():
    print("[8] _safe_write_path 写路径防护（事故回归测试）")
    inside = os.path.join(run_w2.OUT_DIR, "render_pages", "1", "page_1.png")
    run_w2._safe_write_path(inside)  # 不抛即可通过
    check(True, "out/ 内路径通过: %s" % inside)

    raised = False
    try:
        run_w2._safe_write_path(r"G:\投标docx库\_guard_should_raise.docx")
    except AssertionError:
        raised = True
    check(raised, "语料路径必须 raise AssertionError: G:\\投标docx库\\...")

    # 前缀相邻但仍在 out/ 外（考验 out_root + os.sep 判定）
    sibling = os.path.abspath(run_w2.OUT_DIR) + "_evil\\x.json"
    raised2 = False
    try:
        run_w2._safe_write_path(sibling)
    except AssertionError:
        raised2 = True
    check(raised2, "out_root 前缀相邻路径也必须 raise: %s" % sibling)


def test_verdict_empty_doc():
    print("[9] _verdict 空文档判 ok（§5 更新）")
    info = {"parse_error": None}
    r_empty = {"rc_info": 0, "rc_text": 0, "rc_render": 0, "text": "", "truth_norm": 0}
    check(run_w2._verdict(r_empty, info, False) == "ok",
          "truth_norm==0 且空 text → ok（空文档合法对齐）")
    r_missing = {"rc_info": 0, "rc_text": 0, "rc_render": 0, "text": "", "truth_norm": 10}
    check(run_w2._verdict(r_missing, info, False) == "fail_text",
          "truth_norm>0 且空 text → fail_text")
    r_nonempty = {"rc_info": 0, "rc_text": 0, "rc_render": 0, "text": "abc", "truth_norm": 0}
    check(run_w2._verdict(r_nonempty, info, False) == "ok",
          "有输出 → ok")
    r_badtext = {"rc_info": 0, "rc_text": 1, "rc_render": 0, "text": "abc", "truth_norm": 0}
    check(run_w2._verdict(r_badtext, info, False) == "fail_text",
          "rc_text!=0 → fail_text")


def main():
    setup_stdout()
    os.makedirs(FIX_DIR, exist_ok=True)
    try:
        p1 = _make_plain(os.path.join(FIX_DIR, "f1_plain.docx"))
        p2 = _make_table(os.path.join(FIX_DIR, "f2_table.docx"))
        p3 = _make_fullwidth(os.path.join(FIX_DIR, "f3_fullwidth.docx"))
        p4 = _make_del_fixture(os.path.join(FIX_DIR, "f4_del.docx"))

        test_truth_plain(p1)
        test_truth_table(p2)
        test_del(p4)
        test_norm(p3)
        test_metrics()
        a, b, c, d = test_sampling()
        test_pool_counts(a, b, c, d)
        test_write_guard()
        test_verdict_empty_doc()
    except AssertionError as e:
        print("SELFTEST FAIL: %s" % e)
        return 1
    except Exception as e:
        import traceback
        traceback.print_exc()
        print("SELFTEST ERROR: %s: %s" % (type(e).__name__, e))
        return 1

    print("SELFTEST OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
