#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""md2docx.py —— Markdown → 投标 DOCX 转换器（v3.0，解析层换 Python-Markdown）

__version__ = "3.0"

===============================================================================
v3.0 变更说明（2026-10-20，T10 规格书 v2）
===============================================================================
1. **语法层换成熟解析器**（终结"改一处坏一处"）：
   markdown（Python-Markdown）+ stdlib html.parser 建节点树 → 自研 renderer 映射 docx。
   旧版逐行正则硬解析（表格/加粗/图片各写一套正则）已全部移除。
2. **标书层（封面/TOC/暗标/样式族/图号/表格样式）自研保留**，全部 w: XML 直写。
3. 新增能力：
   - A0 A4 页面 XML 直写（11906×16838），--margin 边距参数；--anonymous 清 core/app 属性
   - A1 h1-h3/p/ul/ol/li(嵌套)/table/blockquote/hr/pre/strong/em/code/del/a/img/br
        + 任务列表 ☐、`\\pagebreak` 分页、疑似表格告警、`\\|` 转义、全角归一化
   - A2 表格默认全框线（--table-style three-line 可切），列宽内容权重 + fixed 三件套，
        tblHeader 跨页重复、cantSplit、连续表正确分表、单元格内图片
   - A3 图片等比缩放（≤14.6cm × ≤20cm，仅缩小不放大）、图{章号}-{序号} 题注、
        缺图 collect-all-then-fail（rc=1）、图段/题注段非 exact 行距
   - A4 样式族 Normal/BidBody/BidTableCell/BidFigure/BidCaption/Heading1-3，
        正文 22 磅 exact（w:line=440 lineRule=exact）+ firstLineChars=200
4. **兼容**：旧 CLI `md --project --bidder -o --anonymous` 完全不变，仅新增可选参数。
5. 禁用项：python-docx 的 `line_spacing` 属性（house 铁律，EMU 值不稳致黑块）；
   pip install / pandoc / Node 一律不调用（依赖仅 stdlib + python-docx + markdown）。

v3.0.1（2026-10-21，T10 返工单 R1；零回退，仅加固）：
   a. 畸形 HTML：转换前标签配平检查（不配平 → 告警"未闭合…可能吞并后续内容"）；
      转换后兜底——叶子块/裸文本节点若残留 `**`/`##`/md 链接/`~~` → 告警。
   b. 畸形表超列：markdown tables 扩展按表头列数静默截断多余单元格，转换前先逐行告警。
   c. 列宽：单列内容权重封顶 46（format-lessons 2.2），防超长工程量描述独占版心。
   d. 任务列表：`- [ ] 文本` 渲染为 "☐ 文本"（不再叠加项目符号 "•"），与规格 A1 一致。

用法:
    python md2docx.py 方案.md --project "XX项目" --bidder "XX公司" -o 投标文件.docx

新增可选参数:
    --body-font 仿宋          正文字体（启动检测 C:\\Windows\\Fonts，降级 仿宋_GB2312→仿宋→宋体）
    --line-spacing {22,28}    正文行距磅值（默认 22 → w:line=440 exact）
    --table-style {grid,three-line}  表格样式（默认 grid 全框线）
    --margin top,bottom,left,right   页边距 cm（默认 2.54,2.54,3.17,3.17）
    --download-images         URL 图片下载到临时目录（缺省：URL 图片记缺图并 rc=1）
    --no-cover                不生成封面与目录页
    --verbose                 打印解析/渲染统计

退出码:
    0 正常；1 存在缺失图片（清单已打印）；2 依赖缺失/输入文件不存在
"""

__version__ = "3.0.1"
__all__ = [
    "convert", "ConversionResult", "LEAK_PATTERNS", "INLINE_RESIDUE_KINDS",
    "scan_leaks", "scan_unclosed_html", "warn_unconverted_blocks",
    "resolve_body_font", "font_file_present", "main",
]

import argparse
import os
import re
import shutil
import sys
import tempfile
import urllib.parse
import urllib.request
import zipfile
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

try:
    import markdown
except ImportError:  # pragma: no cover - 环境缺依赖
    print("ERROR: markdown 未安装（Python-Markdown 为规格指定解析器）。", file=sys.stderr)
    sys.exit(2)

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.image.image import Image as DocxImage
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover - 环境缺依赖
    print("ERROR: python-docx 未安装。pip install python-docx", file=sys.stderr)
    sys.exit(2)


# =============================================================================
# 常量
# =============================================================================

TWIP_PER_CM = 567.0             # 1cm = 567 twip（Word 实测口径，见 format-lessons 2.2）
PAGE_W_TW = 11906               # A4 宽 21.0cm
PAGE_H_TW = 16838               # A4 高 29.7cm
DEFAULT_MARGINS = (2.54, 2.54, 3.17, 3.17)   # 上 下 左 右 cm
MAX_IMG_W_CM = 14.6             # 版心上限（A4 21 - 3.17*2 = 14.66）
MAX_IMG_H_CM = 20.0             # 单图高度上限
HEADING_SIZES = {1: 16.0, 2: 14.0, 3: 12.0}
CODE_ASCII_FONT = "Consolas"
EN_FONT = "Times New Roman"
HEI_FONT = "黑体"
PAGEBREAK_SENTINEL = "\ue000"   # markdown 前处理：\pagebreak → 私有区字符
SUSPICIOUS_TABLE_MSG = "疑似表格未识别，请检查头尾竖线"
COL_WEIGHT_CAP = 46.0           # 单列内容权重封顶（format-lessons 2.2：防超长列独占版心）
TABLE_OVERCOL_MSG = "表格第 {row} 行超列，已截断 {n} 列（表头 {head} 列）"
HTML_UNCLOSED_MSG = "检测到未闭合 HTML 标签，可能吞并后续内容"
UNCONVERTED_INLINE_MSG = "检测到未转换的 markdown 行内标记（疑似被 HTML 块吞并）"


# 字体降级链（值为 C:\Windows\Fonts 下的候选文件名，全部小写比对）
FONT_FILES = {
    "仿宋_GB2312": ["fzfsk.ttf", "fzfsjw.ttf", "simfang_gb2312.ttf", "fsgb2312.ttf"],
    "仿宋": ["simfang.ttf"],
    "宋体": ["simsun.ttc", "simsun.ttf"],
    "黑体": ["simhei.ttf"],
    "楷体": ["simkai.ttf"],
    "微软雅黑": ["msyh.ttc"],
    "Times New Roman": ["times.ttf", "timesbd.ttf"],
    "Consolas": ["consola.ttf"],
}
FONT_FALLBACK_CHAIN = ["仿宋_GB2312", "仿宋", "宋体"]

# 11 类语法泄漏标记（A1 断言用；selftest 复用同一份定义）
LEAK_PATTERNS = [
    ("double_asterisk", re.compile(r"\*\*")),
    ("single_asterisk", re.compile(r"(?<=\S)\*(?=\S)")),
    ("hash_heading", re.compile(r"(?:^|(?<=\s))#{1,6}\s*\S", re.M)),
    ("pipe", re.compile(r"\|")),
    ("backtick", re.compile(r"`")),
    ("tilde_strike", re.compile(r"~~")),
    ("md_link", re.compile(r"!?\[[^\]\n]*\]\([^)\n]*\)")),
    ("html_entity", re.compile(r"&(?:[a-zA-Z][a-zA-Z0-9]{1,10}|#\d{1,6}|#x[0-9a-fA-F]{1,6});")),
    ("backslash_escape", re.compile(r"\\[*_\[\]`|#~>()-]")),
    ("task_list_marker", re.compile(r"\[[ xX]\]")),
    ("bare_separator_row", re.compile(r"^\s*\|?[\s:|-]*-{3,}[\s:|-]*\|?\s*$", re.M)),
]


def scan_leaks(text):
    """返回 [(类别, 命中片段), ...]；空列表表示无泄漏。"""
    hits = []
    for name, pat in LEAK_PATTERNS:
        for m in pat.finditer(text or ""):
            hits.append((name, m.group(0)))
    return hits


# 「解析后兜底」用的行内标记子集：只挑正文绝不该出现的形态（pipe 转义、`
# 行内代码` 等属合法内容，故不纳入），命中即说明被 HTML 块吞并、原文泄漏。
# 注意：task_list_marker（`[ ]`）不在其中——任务列表是渲染期才转 ☐ 的，
# 解析树里 `[ ]` 属正常中间态（见 Converter._normalize_task_item）。
INLINE_RESIDUE_KINDS = (
    "double_asterisk", "single_asterisk", "hash_heading",
    "md_link", "tilde_strike",
)


# =============================================================================
# A1 畸形 HTML 检测：① 转换前标签配平 ② 转换后行内标记残留兜底
# =============================================================================

# 只认「已知 HTML 元素名」：避免把 `<https://x>` 之类自动链接误判成标签
_HTML_TAGS = {
    "a", "abbr", "address", "article", "aside", "b", "bdi", "bdo", "blockquote",
    "br", "caption", "center", "cite", "code", "col", "colgroup", "dd", "del",
    "details", "dfn", "div", "dl", "dt", "em", "figcaption", "figure", "font",
    "footer", "h1", "h2", "h3", "h4", "h5", "h6", "head", "header", "hr", "html",
    "i", "img", "input", "ins", "kbd", "label", "li", "main", "mark", "nav", "ol",
    "p", "pre", "q", "rp", "rt", "ruby", "s", "samp", "section", "small", "span",
    "strike", "strong", "sub", "summary", "sup", "table", "tbody", "td", "tfoot",
    "th", "thead", "tr", "tt", "u", "ul", "var", "wbr",
}
_HTML_VOID = {"br", "hr", "img", "input", "wbr", "col", "area", "base", "embed",
              "link", "meta", "param", "source"}
_HTML_TAG_RE = re.compile(
    r"<\s*(/?)\s*([a-zA-Z][a-zA-Z0-9]*)((?:[^<>\"']|\"[^\"]*\"|'[^']*')*?)(/?)\s*>"
)
# 代码块/行内代码/HTML 注释里的尖括号不参与配平
_HTML_MASK_RE = re.compile(r"```.*?```|~~~.*?~~~|<!--.*?-->|`[^`\n]*`", re.S)


def _mask_spans(text, pattern):
    """把匹配片段替换为等长空白（保留换行），使行号不变。"""
    return pattern.sub(lambda m: re.sub(r"[^\n]", " ", m.group(0)), text or "")


def scan_unclosed_html(text):
    """简易标签配平检查（栈）。返回未闭合标签 [(tag, 行号), ...]。

    跳过：代码围栏/行内代码/注释、void 元素、自闭合标签、非已知 HTML 元素名。
    """
    masked = _mask_spans(text, _HTML_MASK_RE)
    stack = []
    for m in _HTML_TAG_RE.finditer(masked):
        closing, tag, selfclose = m.group(1), m.group(2).lower(), m.group(4)
        if tag not in _HTML_TAGS or tag in _HTML_VOID or selfclose:
            continue
        if closing:
            for i in range(len(stack) - 1, -1, -1):
                if stack[i][0] == tag:
                    del stack[i:]
                    break
        else:
            stack.append((tag, masked.count("\n", 0, m.start()) + 1))
    return stack


# 渲染后的叶子块：这些块里的文本若残留 md 行内标记，说明解析器把它们当纯文本了
_LEAF_BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "td", "th"}


def warn_unconverted_blocks(root, warnings, limit=3):
    """② 解析后兜底：块内/裸文本节点残留未转换 md 行内标记 → 告警（复用 LEAK_PATTERNS）。

    正常解析时 `**`/`##`/`[a](b)` 早已变成元素节点；它们以字面字符留在文本里，
    只可能是"HTML 块吞并后原文整段当纯文本"这一类畸形。
    """
    found = []

    def check(text, tag):
        for kind, frag in scan_leaks(text):
            if kind in INLINE_RESIDUE_KINDS:
                found.append((kind, frag, tag))
                return

    def walk(node):
        for child in node.children:
            if child.tag in ("pre", "code"):
                continue                               # 代码内容原样保留，不算泄漏
            if child.tag == "#text":
                if child.data.strip():
                    check(child.data, "#text")
            elif child.tag in _LEAF_BLOCK_TAGS:
                check(child.flatten(), child.tag)
            else:
                walk(child)

    walk(root)
    for kind, frag, tag in found[:limit]:
        label = tag if tag == "#text" else f"<{tag}>"
        warnings.append(f"{UNCONVERTED_INLINE_MSG}：{label} 内残留 {kind} {frag!r}")
    if len(found) > limit:
        warnings.append(f"{UNCONVERTED_INLINE_MSG}：另有 {len(found) - limit} 处")
    return found


# =============================================================================
# 字体检测（A4）
# =============================================================================

def _fonts_dir():
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot") or r"C:\Windows"
    p = Path(windir) / "Fonts"
    return p if p.is_dir() else None


def font_file_present(name):
    """字体名 → 本机 Fonts 目录是否有对应文件。未收录映射返回 None（无法验证）。"""
    key = name.strip()
    if key not in FONT_FILES:
        return None
    fdir = _fonts_dir()
    if fdir is None:
        return False
    have = set()
    try:
        for item in fdir.iterdir():
            have.add(item.name.lower())
    except OSError:                                  # 目录不可读
        return False
    return any(f in have for f in FONT_FILES[key])


def resolve_body_font(requested):
    """降级链 仿宋_GB2312 → 仿宋 → 宋体；返回 (实际字体名, 说明)。"""
    req = (requested or "仿宋").strip()
    chain = [req] + [f for f in FONT_FALLBACK_CHAIN if f != req]
    for cand in chain:
        ok = font_file_present(cand)
        if ok is True:
            note = f"{cand}（C:\\Windows\\Fonts 命中 {FONT_FILES.get(cand, ['?'])[0]}）"
            return cand, note
        if ok is None:                               # 用户自定义字体，无法验证文件名
            return cand, f"{cand}（未收录字体映射，按用户指定使用）"
    return "宋体", "宋体（降级链全部未命中，回落宋体）"


# =============================================================================
# A1 预归一化 + 疑似表格修复
# =============================================================================

_FULLWIDTH_MAP = {"＃": "#", "＊": "*", "｜": "|"}
# 「##标题」→「## 标题」：第二个捕获组必须是非空白非 # 字符，避免 #{1,6} 回溯拆散 '##'
_HEADING_SPACE_RE = re.compile(r"^(#{1,6})([^\s#])")
_SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")
_FENCE_RE = re.compile(r"^\s*(`{3,}|~{3,})")
# Markdown 链接/图片目标：`![alt](目标 "title")`、`[文本](目标)`
_LINK_TARGET_RE = re.compile(r"(!?\[[^\]\n]*\]\()([^)\n]*)(\))")


def _protect_paths(text):
    """链接/图片目标里的反斜杠加倍——否则 markdown 会把 `\\.qidi` 当转义吃掉，
    Windows 绝对路径（G:\\...）会被破坏。"""
    def repl(m):
        target = m.group(2)
        if "\\" in target:
            target = target.replace("\\", "\\\\")
        return m.group(1) + target + m.group(3)
    return _LINK_TARGET_RE.sub(repl, text)


def normalize_syntax(text):
    """全角语法字符归一化 + 行首 '#标题' → '# 标题' + \\pagebreak → 分页哨兵。"""
    for fw, hw in _FULLWIDTH_MAP.items():
        text = text.replace(fw, hw)
    text = re.sub(r"\\pagebreak\b", PAGEBREAK_SENTINEL, text)
    out = []
    for line in text.split("\n"):
        out.append(_HEADING_SPACE_RE.sub(lambda m: m.group(1) + " " + m.group(2),
                                         line, count=1))
    return "\n".join(out)


def split_row_cells(line):
    """按未转义 '|' 切单元格；自动去头尾竖线。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells, cur, i = [], "", 0
    while i < len(s):
        ch = s[i]
        if ch == "\\" and i + 1 < len(s) and s[i + 1] == "|":
            cur += "|"
            i += 2
            continue
        if ch == "|":
            cells.append(cur)
            cur = ""
        else:
            cur += ch
        i += 1
    cells.append(cur)
    return [c.strip() for c in cells]


def _parses_as_table(chunk):
    """用真实解析器做判定：该片段能否被 tables 扩展识别为表格。"""
    md = _new_markdown()
    return "<table>" in md.convert(chunk)


def repair_suspicious_tables(text, warnings):
    """未被 tables 扩展识别的"竖线块"→ 告警 + 可读文本（删分隔行，不泄漏 '|'）。"""
    lines = text.split("\n")
    out, i = [], 0
    in_fence = False
    while i < len(lines):
        line = lines[i]
        if _FENCE_RE.match(line):
            in_fence = not in_fence
            out.append(line)
            i += 1
            continue
        pipeish = (not in_fence) and ("|" in line) and not line[:1].isspace()
        if not pipeish:
            out.append(line)
            i += 1
            continue
        j = i
        while j < len(lines) and lines[j].strip() and "|" in lines[j] and not lines[j][:1].isspace():
            j += 1
        run = lines[i:j]
        if len(run) < 2 or _parses_as_table("\n".join(run)):
            out.extend(run)
            i = j
            continue
        # —— 疑似表格：清洗为可读段落，逐行独立成段 ——
        warnings.append(f"{SUSPICIOUS_TABLE_MSG}（第 {i + 1} 行起，共 {len(run)} 行）")
        kept = 0
        for raw in run:
            cells = split_row_cells(raw)
            real = [c for c in cells if c]
            if real and all(_SEP_CELL_RE.match(c) for c in real):
                continue                                  # 纯分隔行，删除
            text_line = "　".join(c.replace("|", "｜") for c in cells if c)
            if text_line.strip():
                if kept:
                    out.append("")
                out.append(text_line)
                kept += 1
        i = j
    return "\n".join(out)


def warn_over_column_tables(text, warnings):
    """畸形表超列告警：markdown tables 扩展按表头列数**静默截断**多余单元格，
    截断发生在解析器内部（HTML 里已看不到多出的格），故必须在转换前逐行告警。"""
    lines = text.split("\n")
    i, in_fence = 0, False
    while i < len(lines):
        line = lines[i]
        if _FENCE_RE.match(line):
            in_fence = not in_fence
            i += 1
            continue
        pipeish = (not in_fence) and ("|" in line) and not line[:1].isspace()
        if not pipeish:
            i += 1
            continue
        j = i
        while j < len(lines) and lines[j].strip() and "|" in lines[j] and not lines[j][:1].isspace():
            j += 1
        run = lines[i:j]
        if len(run) >= 2 and _parses_as_table("\n".join(run)):
            head = len(split_row_cells(run[0]))
            for k in range(1, len(run)):
                cells = split_row_cells(run[k])
                real = [c for c in cells if c]
                if real and all(_SEP_CELL_RE.match(c) for c in real):
                    continue                          # 分隔行不计列数
                extra = len(cells) - head
                if head and extra > 0:
                    warnings.append(TABLE_OVERCOL_MSG.format(
                        row=i + k + 1, n=extra, head=head))
        i = j


def prepare_markdown(text, warnings):
    """预归一化 → 畸形检测 → 疑似表格修复，返回可交给 Python-Markdown 的文本。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    unclosed = scan_unclosed_html(text)
    if unclosed:
        tag_list = "、".join(sorted({f"<{t}>" for t, _ in unclosed}))
        warnings.append(f"{HTML_UNCLOSED_MSG}：{tag_list}（第 {unclosed[0][1]} 行起）")
    text = normalize_syntax(text)
    text = _protect_paths(text)
    text = repair_suspicious_tables(text, warnings)
    warn_over_column_tables(text, warnings)
    return text


# =============================================================================
# A1 解析：markdown → HTML → stdlib html.parser 节点树
# =============================================================================

try:  # markdown 未提供 ~~ 删除线（核心无 del 扩展），自建极简 InlineProcessor
    from markdown.extensions import Extension as _MdExtension
    from markdown.inlinepatterns import InlineProcessor as _InlineProcessor
    import xml.etree.ElementTree as _etree

    class _DelInlineProcessor(_InlineProcessor):
        def handleMatch(self, m, data):
            el = _etree.Element("del")
            el.text = m.group(1)
            return el, m.start(0), m.end(0)

    class StrikethroughExtension(_MdExtension):
        """`~~文本~~` → <del>文本</del>（A1 要求渲染删除线）。

        Python-Markdown 核心无 strikethrough，故补一个 15 行内联扩展；
        规格指定的 tables / fenced_code / sane_lists 三个扩展保持原样。
        """

        def extendMarkdown(self, md):
            md.inlinePatterns.register(
                _DelInlineProcessor(r"~~(?=\S)(.+?)(?<=\S)~~", md), "qidi_del", 175
            )
except Exception:  # pragma: no cover - markdown 版本差异兜底
    StrikethroughExtension = None


def _new_markdown():
    exts = ["tables", "fenced_code", "sane_lists"]
    if StrikethroughExtension is not None:
        exts.append(StrikethroughExtension())
    return markdown.Markdown(extensions=exts)


class Node(object):
    """HTML 节点树节点。tag == '#text' 时 data 为文本。"""

    __slots__ = ("tag", "attrs", "children", "data")

    def __init__(self, tag, attrs=None, data=""):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.children = []
        self.data = data

    def __repr__(self):  # pragma: no cover - 调试用
        return f"<Node {self.tag} {self.attrs} {self.data!r} n={len(self.children)}>"

    def flatten(self, strip_pre=False):
        if self.tag == "#text":
            return self.data
        return "".join(c.flatten(strip_pre) for c in self.children)

    def shallow(self, data):
        return Node(self.tag, self.attrs, data)


_VOID_TAGS = {"br", "hr", "img", "meta", "link", "input", "col", "area", "base",
              "source", "wbr", "param", "embed"}


class _TreeBuilder(HTMLParser):
    def __init__(self):
        HTMLParser.__init__(self, convert_charrefs=True)
        self.root = Node("#root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag.lower(), attrs)
        self.stack[-1].children.append(node)
        if tag.lower() not in _VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag.lower(), attrs))

    def handle_endtag(self, tag):
        tag = tag.lower()
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if data:
            self.stack[-1].children.append(Node("#text", data=data))

    def handle_comment(self, data):
        pass

    def handle_decl(self, decl):
        pass

    def handle_pi(self, data):
        pass


def html_to_tree(html):
    # 分页哨兵 → 元素节点，交给 renderer 切成两个段落 + 一个分页符
    html = html.replace(PAGEBREAK_SENTINEL, "<qidi-pagebreak></qidi-pagebreak>")
    tb = _TreeBuilder()
    tb.feed(html)
    tb.close()
    return tb.root


_CJK_RE = re.compile(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef]")


def clean_text(s):
    """软换行/&nbsp; 归一化：中文之间软换行直接相连，其余变空格。"""
    if not s:
        return s
    s = s.replace("\xa0", " ")
    s = re.sub(r"[ \t]*\n[ \t]*", "\x01", s)
    s = re.sub(r"(?<=%s)\x01(?=%s)" % (_CJK_RE.pattern, _CJK_RE.pattern), "", s)
    return s.replace("\x01", " ")


# =============================================================================
# docx XML 小工具
# =============================================================================

_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def set_spacing(pPr, line_twips=None, rule="exact", before=None, after=None):
    """行距 XML 直写（禁止 python-docx line_spacing 属性）。"""
    sp = pPr.get_or_add_spacing()
    if line_twips is not None:
        sp.set(qn("w:line"), str(int(line_twips)))
    if rule:
        sp.set(qn("w:lineRule"), rule)
    if before is not None:
        sp.set(qn("w:before"), str(int(before)))
    if after is not None:
        sp.set(qn("w:after"), str(int(after)))


def set_ind(pPr, first_line_chars=None, first_line=None, left=None, hanging=None):
    ind = pPr.get_or_add_ind()
    if first_line_chars is not None:
        ind.set(qn("w:firstLineChars"), str(int(first_line_chars)))
        if int(first_line_chars) == 0:
            ind.set(qn("w:firstLine"), "0")
    if first_line is not None:
        ind.set(qn("w:firstLine"), str(int(first_line)))
    if left is not None:
        ind.set(qn("w:left"), str(int(left)))
    if hanging is not None:
        ind.set(qn("w:hanging"), str(int(hanging)))
    return ind


def set_style_fonts(style_el, east_asia=None, ascii_font=None, size_pt=None,
                    bold=None, color_hex=None):
    rPr = style_el.get_or_add_rPr()
    rF = rPr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rF.get(qn(attr)) is not None:
            del rF.attrib[qn(attr)]
    if east_asia:
        rF.set(qn("w:eastAsia"), east_asia)
    if ascii_font:
        rF.set(qn("w:ascii"), ascii_font)
        rF.set(qn("w:hAnsi"), ascii_font)
    if size_pt:
        half = str(int(round(size_pt * 2)))
        rPr.get_or_add_sz().set(qn("w:val"), half)
        szcs = rPr.find(qn("w:szCs"))
        if szcs is not None:                          # 已存在则同步，避免二号槽不一致
            szcs.set(qn("w:val"), half)
    if bold is not None:
        (rPr.get_or_add_b() if bold else rPr.get_or_add_b()).set(qn("w:val"), "1" if bold else "0")
    if color_hex:
        col = rPr.get_or_add_color()
        col.set(qn("w:val"), color_hex)
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if col.get(qn(attr)) is not None:
                del col.attrib[qn(attr)]


def set_table_borders(tblPr, edges, val="single", sz=4, color="000000"):
    """表级 tblBorders；edges 为 6 边名子集。"""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription",
    )
    return borders


def set_cell_borders(cell, **edges):
    """单元格边框（三线表用）；edges: top=dict(sz=..,val=..)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcB = tcPr.find(qn("w:tcBorders"))
    if tcB is None:
        tcB = OxmlElement("w:tcBorders")
        tcPr.append(tcB)
    for edge, spec in edges.items():
        el = tcB.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcB.append(el)
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", 4)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "000000"))


def set_cell_width(cell, twips):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.insert_element_before(tcW, "w:gridSpan", "w:hMerge", "w:vMerge",
                                   "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar",
                                   "w:textDirection", "w:tcFitText", "w:vAlign",
                                   "w:hideMark")
    tcW.set(qn("w:w"), str(int(twips)))
    tcW.set(qn("w:type"), "dxa")


def distribute_widths(weights, total_twips, min_twips):
    """按权重分配列宽，总和严格等于 total_twips，每列不低于 min_twips。"""
    n = len(weights)
    if n == 0:
        return []
    w = [max(float(x), 0.1) for x in weights]
    total_twips = int(total_twips)
    if min_twips * n >= total_twips:                 # 列太多，退化为均分
        base = total_twips // n
        out = [base] * n
        out[-1] += total_twips - sum(out)
        return out
    fixed = [None] * n
    remaining = list(range(n))
    avail = total_twips
    while remaining:
        s = sum(w[i] for i in remaining)
        changed = False
        for i in list(remaining):
            share = avail * w[i] / s if s else 0
            if share < min_twips:
                fixed[i] = min_twips
                avail -= min_twips
                remaining.remove(i)
                changed = True
        if not changed:
            s = sum(w[i] for i in remaining)
            for i in remaining:
                fixed[i] = int(round(avail * w[i] / s)) if s else 0
            break
    out = [int(fixed[i]) for i in range(n)]
    drift = total_twips - sum(out)
    if drift:                                        # 尾差补到最宽列
        idx = max(range(n), key=lambda k: out[k])
        out[idx] += drift
    return out


def _text_width(text):
    return sum(1.0 if ord(ch) > 127 else 0.55 for ch in text or "")


# =============================================================================
# 转换结果
# =============================================================================

class ConversionResult(object):
    def __init__(self):
        self.warnings = []
        self.missing_images = []
        self.tables = 0
        self.figures = 0
        self.paragraphs = 0
        self.list_items = 0
        self.downloaded = []
        self.body_font = ""
        self.out_path = ""

    @property
    def rc(self):
        return 1 if self.missing_images else 0

    def summary(self):
        return (f"tables={self.tables} figures={self.figures} "
                f"paragraphs={self.paragraphs} list_items={self.list_items} "
                f"warnings={len(self.warnings)} missing_images={len(self.missing_images)}")


# =============================================================================
# 主转换器
# =============================================================================

class Converter(object):
    def __init__(self, md_path, out_path, project="", bidder="", anonymous=False,
                 body_font="仿宋", line_spacing=22, table_style="grid",
                 margins=DEFAULT_MARGINS, download_images=False, cover=True,
                 verbose=False):
        self.md_path = Path(md_path)
        self.out_path = Path(out_path)
        self.project = project or "投标文件"
        self.bidder = bidder or ""
        self.anonymous = bool(anonymous)
        self.requested_font = body_font
        self.line_spacing = int(line_spacing)
        self.table_style = table_style
        self.margins = tuple(float(x) for x in margins)
        self.download_images = bool(download_images)
        self.cover = bool(cover)
        self.verbose = bool(verbose)

        self.result = ConversionResult()
        self.warnings = self.result.warnings
        self.missing_images = self.result.missing_images

        self.doc = None
        self.tmp_dirs = []
        self.h1_count = 0
        self.fig_seq = {}
        self.body_font = "宋体"
        self.body_font_note = ""
        self.line_twips = self.line_spacing * 20
        left_tw = int(round(self.margins[2] * TWIP_PER_CM))
        right_tw = int(round(self.margins[3] * TWIP_PER_CM))
        self.content_tw = PAGE_W_TW - left_tw - right_tw      # 与 pgMar XML 严格一致

    # ---------------------------------------------------------------- 样式族
    def _ensure_style(self, name, base="Normal"):
        try:
            return self.doc.styles[name]
        except KeyError:
            st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = self.doc.styles[base]
            return st

    def build_styles(self):
        doc = self.doc
        # Normal：仅三槽字体 + 12pt（不设行距/缩进）
        normal = doc.styles["Normal"]
        set_style_fonts(normal.element, east_asia=self.body_font,
                        ascii_font=EN_FONT, size_pt=12)

        # BidBody：22 磅 exact + 首行缩进 2 字符
        body = self._ensure_style("BidBody")
        pPr = body.element.get_or_add_pPr()
        set_spacing(pPr, line_twips=self.line_twips, rule="exact", before=0, after=0)
        set_ind(pPr, first_line_chars=200)

        # BidTableCell：五号、单倍、无缩进
        cellst = self._ensure_style("BidTableCell")
        pPr = cellst.element.get_or_add_pPr()
        set_spacing(pPr, line_twips=240, rule="auto", before=0, after=0)
        set_ind(pPr, first_line_chars=0, first_line=0, left=0)
        set_style_fonts(cellst.element, east_asia=self.body_font,
                        ascii_font=EN_FONT, size_pt=10.5)

        # BidFigure：居中、无 exact、段前后 6pt
        fig = self._ensure_style("BidFigure")
        pPr = fig.element.get_or_add_pPr()
        set_spacing(pPr, line_twips=240, rule="auto", before=120, after=120)
        set_ind(pPr, first_line_chars=0, first_line=0)
        pPr.get_or_add_jc().set(qn("w:val"), "center")

        # BidCaption：居中 10.5pt 黑体、无 exact
        cap = self._ensure_style("BidCaption")
        pPr = cap.element.get_or_add_pPr()
        set_spacing(pPr, line_twips=240, rule="auto", before=0, after=120)
        set_ind(pPr, first_line_chars=0, first_line=0)
        pPr.get_or_add_jc().set(qn("w:val"), "center")
        set_style_fonts(cap.element, east_asia=HEI_FONT, ascii_font=HEI_FONT,
                        size_pt=10.5, color_hex="000000")

        # Heading 1/2/3：样式级覆盖 黑体 + 黑色 + 16/14/12pt
        for lvl in (1, 2, 3):
            st = doc.styles[f"Heading {lvl}"]
            set_style_fonts(st.element, east_asia=HEI_FONT, ascii_font=HEI_FONT,
                            size_pt=HEADING_SIZES[lvl], color_hex="000000")
            pPr = st.element.get_or_add_pPr()
            if self.line_spacing != 22:              # 28 磅用于标题（规格 A4）
                set_spacing(pPr, line_twips=self.line_twips, rule="exact")
            else:
                set_spacing(pPr, line_twips=240, rule="auto")

    # ---------------------------------------------------------------- 页面
    def setup_page(self):
        sectPr = self.doc.sections[0]._sectPr
        pgSz = sectPr.find(qn("w:pgSz"))
        if pgSz is None:
            pgSz = OxmlElement("w:pgSz")
            sectPr.insert(0, pgSz)
        pgSz.set(qn("w:w"), str(PAGE_W_TW))
        pgSz.set(qn("w:h"), str(PAGE_H_TW))
        pgMar = sectPr.find(qn("w:pgMar"))
        if pgMar is None:
            pgMar = OxmlElement("w:pgMar")
            sectPr.append(pgMar)
        top, bottom, left, right = self.margins
        pgMar.set(qn("w:top"), str(int(round(top * TWIP_PER_CM))))
        pgMar.set(qn("w:bottom"), str(int(round(bottom * TWIP_PER_CM))))
        pgMar.set(qn("w:left"), str(int(round(left * TWIP_PER_CM))))
        pgMar.set(qn("w:right"), str(int(round(right * TWIP_PER_CM))))
        pgMar.set(qn("w:header"), "720")
        pgMar.set(qn("w:footer"), "720")
        pgMar.set(qn("w:gutter"), "0")

    # ---------------------------------------------------------------- run
    def add_run(self, paragraph, text, bold=False, italic=False, strike=False,
                code=False, size=12.0, east_asia=None, color_hex=None):
        if text == "":
            return None
        r = paragraph.add_run(text)
        if code:
            ea = east_asia or self.body_font
            ascii_font = CODE_ASCII_FONT
        else:
            ea = east_asia or self.body_font
            ascii_font = EN_FONT
        rPr = r._element.get_or_add_rPr()
        rF = rPr.get_or_add_rFonts()
        rF.set(qn("w:ascii"), ascii_font)
        rF.set(qn("w:hAnsi"), ascii_font)
        rF.set(qn("w:eastAsia"), ea)
        r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if italic:
            r.font.italic = True
        if strike:
            r.font.strike = True
        if color_hex:
            r.font.color.rgb = RGBColor.from_string(color_hex)
        return r

    # ---------------------------------------------------------------- 封面/TOC
    def _cover_par(self, text="", size=14.0, east_asia=None, bold=False):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p._p.get_or_add_pPr(), line_twips=240, rule="auto")  # 防 exact 压大字
        if text:
            self.add_run(p, text, size=size, east_asia=east_asia or self.body_font, bold=bold)
        return p

    def build_cover(self):
        if not self.cover:
            return
        if self.anonymous:
            # 暗标：不出现投标人信息，仅"投标文件"（保持 v2 行为：无封面无目录）
            return
        for _ in range(4):
            self._cover_par()
        self._cover_par(self.project, size=26.0, east_asia=HEI_FONT, bold=True)
        self._cover_par()
        self._cover_par("投 标 文 件", size=18.0, east_asia=HEI_FONT, bold=True)
        self._cover_par()
        if self.bidder:
            self._cover_par(f"投标人：{self.bidder}", size=14.0)
        self._cover_par()
        self._cover_par(f"日期：{date.today().strftime('%Y年%m月%d日')}", size=14.0)
        self.doc.add_page_break()
        self._cover_par("目  录", size=16.0, east_asia=HEI_FONT, bold=True)
        self.add_toc_field()
        self.doc.add_page_break()

    def add_toc_field(self):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        run = p.add_run()
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-3" \\h \\z \\u'
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        run._r.append(b)
        run._r.append(instr)
        run._r.append(e)

    # ---------------------------------------------------------------- 块级渲染
    def body_par(self, style="BidBody", indent_left=None, first_line_chars=None,
                 quote=0):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles[style]
        pPr = p._p.get_or_add_pPr()
        left = indent_left if indent_left is not None else (420 * quote if quote else None)
        if left is not None or first_line_chars is not None:
            set_ind(pPr, first_line_chars=first_line_chars, first_line=0, left=left)
        self.result.paragraphs += 1
        return p

    def render(self, root):
        self.render_blocks(root.children)

    def render_blocks(self, nodes, quote=0):
        for n in nodes:
            tag = n.tag
            if tag == "#text":
                if n.data.strip():
                    self.render_paragraph(Node("p", {}, ""), [n], quote=quote)
                continue
            if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
                self.render_heading(n)
            elif tag == "p":
                self.render_paragraph(n, n.children, quote=quote)
            elif tag == "ul":
                self.render_list(n, ordered=False, level=0, quote=quote)
            elif tag == "ol":
                self.render_list(n, ordered=True, level=0, quote=quote)
            elif tag == "table":
                self.render_table(n)
            elif tag == "blockquote":
                self.render_blocks(n.children, quote=quote + 1)
            elif tag == "hr":
                self.render_hr()
            elif tag == "pre":
                self.render_pre(n)
            elif tag == "qidi-pagebreak":
                self.doc.add_page_break()
            else:                                    # div/span/section 等透明容器
                self.render_blocks(n.children, quote=quote)

    def render_heading(self, node):
        lvl = min(int(node.tag[1]), 3)
        if node.tag == "h1":
            self.h1_count += 1
        p = self.doc.add_paragraph()
        p.style = self.doc.styles[f"Heading {lvl}"]
        self.result.paragraphs += 1
        self.add_inline(p, node.children,
                        {"size": HEADING_SIZES[lvl], "east_asia": HEI_FONT,
                         "bold": True, "color_hex": "000000"})
        self._drop_trailing_blank_runs(p)

    def _split_pagebreak(self, children):
        """按分页哨兵节点切段：[[节点...], None(分页), [节点...]]"""
        segs, cur = [], []
        for c in children:
            if c.tag == "qidi-pagebreak":
                segs.append(cur)
                segs.append(None)
                cur = []
            else:
                cur.append(c)
        segs.append(cur)
        return segs

    def render_paragraph(self, node, children, quote=0):
        for seg in self._split_pagebreak(children):
            if seg is None:
                self.doc.add_page_break()
                continue
            sig = [c for c in seg if not (c.tag == "#text" and not c.data.strip())]
            if not sig:
                continue
            if len(sig) == 1 and sig[0].tag == "img":
                self.render_figure(sig[0])
                continue
            text_peek = "".join(c.flatten() for c in sig)
            if "|" in text_peek and "---" in text_peek:
                self.warnings.append(SUSPICIOUS_TABLE_MSG + "：" + text_peek.strip()[:40])
            p = self.body_par(quote=quote, first_line_chars=0 if quote else None)
            self.add_inline(p, seg, {})
            self._drop_trailing_blank_runs(p)

    def _drop_trailing_blank_runs(self, paragraph):
        for r in reversed(paragraph.runs):
            if r.text and not r.text.strip():
                r._r.getparent().remove(r._r)
            elif r.text:
                break

    def render_list(self, node, ordered, level, quote=0):
        idx = 1
        for li in [c for c in node.children if c.tag == "li"]:
            # li 内可能有 p（松散列表）与嵌套 ul/ol
            inline_children, nested = [], []
            for c in li.children:
                if c.tag in ("ul", "ol"):
                    nested.append(c)
                elif c.tag == "p":
                    inline_children.extend(c.children)
                else:
                    inline_children.append(c)
            # 任务列表：首文本节点 "[ ] xxx" / "[x] xxx" → "☐ xxx" / "☑ xxx"
            # 规格 A1：`- [ ] 文本` → "☐ 文本"（不再叠加项目符号 "•"，否则成 "• ☐"）
            inline_children, is_task = self._normalize_task_item(inline_children, ordered)
            marker = "" if is_task else (f"{idx}." if ordered else "•")
            p = self.body_par(indent_left=420 * level, first_line_chars=0,
                              quote=0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if marker:
                self.add_run(p, marker + " ", size=12.0, bold=False)
            self.add_inline(p, inline_children, {})
            self._drop_trailing_blank_runs(p)
            self.result.list_items += 1
            for sub in nested:
                self.render_list(sub, ordered=(sub.tag == "ol"), level=level + 1, quote=quote)
            idx += 1

    def _normalize_task_item(self, children, ordered):
        """返回 (children, is_task)；is_task=True 时不再加项目符号（规格 A1）。"""
        if ordered or not children:
            return children, False
        first = children[0]
        if first.tag != "#text":
            return children, False
        m = re.match(r"^\[([ xX])\]\s*", first.data)
        if not m:
            return children, False
        mark = "☑ " if m.group(1).lower() == "x" else "☐ "
        out = list(children)
        out[0] = first.shallow(mark + first.data[m.end():])
        return out, True

    def render_blockquote(self, node):
        self.render_blocks(node.children, quote=1)

    def render_hr(self):
        p = self.doc.add_paragraph()
        p.style = self.doc.styles["Normal"]
        pPr = p._p.get_or_add_pPr()
        set_spacing(pPr, line_twips=240, rule="auto")
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        pBdr.append(bottom)
        pPr.insert_element_before(pBdr, *_PPR_AFTER_PBDR)

    def render_pre(self, node):
        code = node
        for c in node.children:
            if c.tag == "code":
                code = c
                break
        text = code.flatten().rstrip("\n")
        p = self.body_par(style="BidTableCell", first_line_chars=0, indent_left=0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lines = text.split("\n")
        for i, line in enumerate(lines):
            if i:
                r = p.add_run()
                r.add_break()
            self.add_run(p, line, code=True, size=10.5)

    # ---------------------------------------------------------------- 行内渲染
    _BOLD = {"strong", "b"}
    _ITALIC = {"em", "i", "cite", "var"}
    _STRIKE = {"del", "s", "strike"}
    _CODE = {"code", "kbd", "samp", "tt"}

    def add_inline(self, paragraph, nodes, fmt, depth=0, in_cell=False,
                   cell_max_cm=None):
        for n in nodes:
            tag = n.tag
            if tag == "#text":
                txt = clean_text(n.data)
                if not txt:
                    continue
                self.add_run(paragraph, txt,
                             bold=fmt.get("bold", False),
                             italic=fmt.get("italic", False),
                             strike=fmt.get("strike", False),
                             code=fmt.get("code", False),
                             size=fmt.get("size", 12.0),
                             east_asia=fmt.get("east_asia"),
                             color_hex=fmt.get("color_hex"))
            elif tag in self._BOLD:
                self._recurse(paragraph, n, fmt, depth, in_cell, cell_max_cm, bold=True)
            elif tag in self._ITALIC:
                self._recurse(paragraph, n, fmt, depth, in_cell, cell_max_cm, italic=True)
            elif tag in self._STRIKE:
                self._recurse(paragraph, n, fmt, depth, in_cell, cell_max_cm, strike=True)
            elif tag in self._CODE:
                sub = dict(fmt)
                sub["code"] = True
                self.add_inline(paragraph, n.children, sub, depth + 1, in_cell,
                                cell_max_cm)
            elif tag == "a":
                href = n.attrs.get("href", "")
                self.add_inline(paragraph, n.children, fmt, depth + 1, in_cell,
                                cell_max_cm)
                if href:
                    self.add_run(paragraph, f"（{href}）",
                                 size=fmt.get("size", 12.0),
                                 east_asia=fmt.get("east_asia"),
                                 color_hex=fmt.get("color_hex"))
            elif tag == "img":
                if depth == 0 and in_cell:
                    self.insert_cell_image(paragraph, n, cell_max_cm)
                else:
                    alt = n.attrs.get("alt", "") or Path(n.attrs.get("src", "")).name
                    self.warnings.append(f"行内图形态不支持，降级为文本：{alt}")
                    self.add_run(paragraph, f"【图：{alt}】", size=fmt.get("size", 12.0),
                                 east_asia=fmt.get("east_asia"))
            elif tag == "br":
                r = paragraph.add_run()
                r.add_break()
            elif tag in ("sup", "sub"):
                self.add_inline(paragraph, n.children, fmt, depth + 1, in_cell,
                                cell_max_cm)
            else:                                    # span/qidi-* 等透明元素
                self.add_inline(paragraph, n.children, fmt, depth + 1, in_cell,
                                cell_max_cm)

    def _recurse(self, paragraph, node, fmt, depth, in_cell, cell_max_cm, **flags):
        sub = dict(fmt)
        sub.update(flags)
        self.add_inline(paragraph, node.children, sub, depth + 1, in_cell, cell_max_cm)

    # ---------------------------------------------------------------- 图片（A3）
    def _resolve_image(self, src, alt):
        """返回本地路径 Path 或 None（记缺图）。"""
        src = (src or "").strip()
        if not src:
            self.missing_images.append(f"（空路径）alt={alt}")
            return None
        is_win_abs = bool(re.match(r"^[a-zA-Z]:[\\/]", src)) or src.startswith("\\\\")
        if (not is_win_abs) and (re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", src)
                                 or src.startswith("//")):
            if src.lower().startswith("data:"):
                self.missing_images.append(f"data URL 不支持：alt={alt}")
                return None
            if not self.download_images:
                self.missing_images.append(f"{src}（URL 图片，未加 --download-images）")
                return None
            try:
                td = Path(tempfile.mkdtemp(prefix="md2docx_img_"))
                self.tmp_dirs.append(td)
                name = Path(urllib.parse.urlparse(src).path).name or "download.png"
                dest = td / name
                req = urllib.request.Request(src, headers={"User-Agent": "md2docx/3.0"})
                with urllib.request.urlopen(req, timeout=30) as resp, open(dest, "wb") as fh:
                    shutil.copyfileobj(resp, fh)
                self.result.downloaded.append(str(dest))
                return dest
            except Exception as exc:                 # 网络/解析失败 → 缺图
                self.missing_images.append(f"{src}（下载失败：{exc}）")
                return None
        p = Path(urllib.parse.unquote(src))
        if not p.is_absolute():
            p = self.md_path.parent / p
        if not p.exists():
            self.missing_images.append(str(p))
            return None
        return p

    def _fit_image(self, path, max_w_cm, max_h_cm):
        """等比缩放（仅缩小不放大），返回 (width_cm or None, height_cm or None)。"""
        try:
            info = DocxImage.from_file(str(path))
            w_cm, h_cm = info.width.cm, info.height.cm
        except Exception as exc:
            self.warnings.append(f"图片尺寸读取失败（{path}）：{exc}")
            return None, None
        if w_cm <= 0 or h_cm <= 0:
            return None, None
        scale = min(1.0, max_w_cm / w_cm, max_h_cm / h_cm)
        if scale >= 1.0:
            return w_cm, h_cm
        return w_cm * scale, h_cm * scale

    def render_figure(self, img_node):
        alt = img_node.attrs.get("alt", "").strip()
        src = img_node.attrs.get("src", "")
        path = self._resolve_image(src, alt)
        chapter = self.h1_count
        if path is None:
            p = self.body_par(style="BidFigure", first_line_chars=0, indent_left=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.add_run(p, f"【缺图：{alt or src}】", size=10.5)
            return
        w_cm, h_cm = self._fit_image(path, MAX_IMG_W_CM, MAX_IMG_H_CM)
        p = self.body_par(style="BidFigure", first_line_chars=0, indent_left=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p._p.get_or_add_pPr(), line_twips=240, rule="auto",
                    before=120, after=120)          # 显式清除 exact，防压图
        run = p.add_run()
        try:
            if w_cm is None:
                run.add_picture(str(path))
            else:
                run.add_picture(str(path), width=Cm(w_cm), height=Cm(h_cm))
        except Exception as exc:
            self.warnings.append(f"图片插入失败（{path}）：{exc}")
            self.missing_images.append(f"{path}（插入失败：{exc}）")
            return
        seq = self.fig_seq.get(chapter, 0) + 1
        self.fig_seq[chapter] = seq
        self.result.figures += 1
        cap = self.body_par(style="BidCaption", first_line_chars=0, indent_left=0)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(cap._p.get_or_add_pPr(), line_twips=240, rule="auto",
                    before=0, after=120)
        label = f"图{chapter}-{seq}"
        self.add_run(cap, f"{label} {alt}".strip(), size=10.5, east_asia=HEI_FONT)

    def insert_cell_image(self, paragraph, img_node, cell_max_cm):
        alt = img_node.attrs.get("alt", "").strip()
        src = img_node.attrs.get("src", "")
        path = self._resolve_image(src, alt)
        if path is None:
            self.add_run(paragraph, f"【图：{alt or src}】", size=10.5)
            return
        cap = min(MAX_IMG_W_CM, (cell_max_cm or MAX_IMG_W_CM) - 0.2)
        w_cm, h_cm = self._fit_image(path, cap, MAX_IMG_H_CM)
        run = paragraph.add_run()
        try:
            if w_cm is None:
                run.add_picture(str(path))
            else:
                run.add_picture(str(path), width=Cm(w_cm), height=Cm(h_cm))
        except Exception as exc:
            self.warnings.append(f"单元格图片插入失败（{path}）：{exc}")
            self.add_run(paragraph, f"【图：{alt or src}】", size=10.5)

    # ---------------------------------------------------------------- 表格（A2）
    @staticmethod
    def _row_cells(tr):
        return [td for td in tr.children if td.tag in ("td", "th")]

    @staticmethod
    def _is_separator_row(tr):
        cells = Converter._row_cells(tr)
        texts = [c.flatten().strip() for c in cells]
        real = [t for t in texts if t]
        return bool(real) and all(_SEP_CELL_RE.match(t) for t in real)

    def render_table(self, node):
        header, body = [], []
        for child in node.children:
            if child.tag == "thead":
                header.extend(tr for tr in child.children if tr.tag == "tr")
            elif child.tag == "tbody":
                body.extend(tr for tr in child.children if tr.tag == "tr")
            elif child.tag == "tr":
                body.append(child)
        if not header and body:
            header = [body.pop(0)]
        if not header:
            self.warnings.append("表格无表头，已跳过：" + node.flatten()[:30])
            return
        # 连续表：tbody 中出现纯分隔行 → 从此处分表（不被吞并、不泄漏 '---'）
        segments = [[]]
        for tr in body:
            if self._is_separator_row(tr):
                segments.append([])
            else:
                segments[-1].append(tr)
        for idx, seg in enumerate(segments):
            if seg or idx == 0:          # 第一段允许只有表头；分隔行造成的空尾段丢弃
                self._build_table(header, seg)

    def _build_table(self, header_rows, body_rows):
        ncols = 0
        for tr in header_rows + body_rows:
            ncols = max(ncols, len(self._row_cells(tr)))
        if ncols == 0:
            return
        table = self.doc.add_table(rows=len(header_rows) + len(body_rows), cols=ncols)
        table.style = self.doc.styles["Table Grid"] if "Table Grid" in [
            s.name for s in self.doc.styles] else None
        table.alignment = 1  # WD_TABLE_ALIGNMENT.CENTER

        # —— 列宽：内容权重（单列封顶 46，防超长工程量列独占版心）→ 总和=版心宽 ——
        weights = []
        for j in range(ncols):
            mx = 0.0
            for tr in header_rows + body_rows:
                cells = self._row_cells(tr)
                if j < len(cells):
                    mx = max(mx, _text_width(cells[j].flatten()))
            weights.append(max(min(mx, COL_WEIGHT_CAP), 3.0))
        widths = distribute_widths(weights, self.content_tw, int(1.6 * TWIP_PER_CM))

        tblPr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook",
                                    "w:tblCaption", "w:tblDescription")
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.insert_element_before(tblW, "w:jc", "w:tblCellSpacing", "w:tblInd",
                                        "w:tblBorders", "w:tblLayout", "w:tblCellMar",
                                        "w:tblLook", "w:tblCaption", "w:tblDescription")
        tblW.set(qn("w:w"), str(sum(widths)))
        tblW.set(qn("w:type"), "dxa")
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(int(w)))

        if self.table_style == "three-line":
            set_table_borders(tblPr, edges=())
        else:
            set_table_borders(tblPr, edges=("top", "left", "bottom", "right",
                                           "insideH", "insideV"), sz=4)

        all_rows = header_rows + body_rows
        for i, tr in enumerate(all_rows):
            is_header = i < len(header_rows)
            is_last = i == len(all_rows) - 1
            row = table.rows[i]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))            # 先 cantSplit
            if is_header and len(body_rows) > 0:
                trPr.append(OxmlElement("w:tblHeader"))        # 后 tblHeader
            cells = self._row_cells(tr)
            for j in range(ncols):
                cell = row.cells[j]
                set_cell_width(cell, widths[j])
                if self.table_style == "three-line":
                    if is_header:
                        set_cell_borders(cell, top={"sz": 12}, bottom={"sz": 4})
                    elif is_last:
                        set_cell_borders(cell, bottom={"sz": 12})
                para = cell.paragraphs[0]
                para.style = self.doc.styles["BidTableCell"]
                # 行超列（单元格数 > 表头列数）在 prepare_markdown 阶段已告警：
                # markdown tables 扩展在解析时就按表头列数截断了多余格。此处只处理
                # 行短于表头的情形（补空格），不留任何静默路径。
                if j < len(cells):
                    src = cells[j]
                    aligned = is_header or (src.tag == "th")
                    if aligned:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.add_inline(para, src.children, {"size": 10.5}, in_cell=True,
                                    cell_max_cm=widths[j] / TWIP_PER_CM)
                    self._drop_trailing_blank_runs(para)
        self.result.tables += 1

    # ---------------------------------------------------------------- 文档属性
    def apply_core_properties(self):
        """A0 暗标：core.xml 的 creator/lastModifiedBy/description/title 清零/脱敏。"""
        cp = self.doc.core_properties
        if self.anonymous:
            cp.author = ""
            cp.last_modified_by = ""
            cp.comments = ""
            cp.title = "投标文件"
            cp.category = ""
            cp.keywords = ""
            cp.subject = ""
        else:
            cp.author = self.bidder or ""
            cp.last_modified_by = self.bidder or ""
            if not cp.title:
                cp.title = self.project
            if "python-docx" in (cp.comments or ""):
                cp.comments = ""

    def scrub_app_properties(self, docx_path):
        """A0 暗标：app.xml 的 Company 置空（python-docx 未暴露该属性，需重写 zip 部件）。"""
        if not self.anonymous:
            return
        self._rewrite_zip_part(docx_path, "docProps/app.xml", self._empty_company)

    @staticmethod
    def _empty_company(xml_text):
        xml_text = re.sub(r"<Company>.*?</Company>", "<Company></Company>",
                          xml_text, flags=re.S)
        if "<Company" not in xml_text and "Properties" in xml_text:
            xml_text = re.sub(r"(<Properties[^>]*>)", r"\1<Company></Company>",
                              xml_text, count=1)
        return xml_text

    @staticmethod
    def _rewrite_zip_part(path, part_name, transform):
        path = Path(path)
        tmp = path.with_name(path.name + ".tmp")
        with zipfile.ZipFile(path, "r") as zin, \
                zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == part_name:
                    data = transform(data.decode("utf-8")).encode("utf-8")
                zout.writestr(item, data)
        tmp.replace(path)

    # ---------------------------------------------------------------- 主流程
    def build(self):
        self.body_font, self.body_font_note = resolve_body_font(self.requested_font)
        print(f"[font] 正文字体: {self.body_font_note}", file=sys.stderr)
        self.result.body_font = self.body_font

        self.doc = Document()
        self.setup_page()
        self.build_styles()
        self.build_cover()

        raw = self.md_path.read_text(encoding="utf-8-sig")
        prepared = prepare_markdown(raw, self.warnings)
        html = _new_markdown().convert(prepared)
        root = html_to_tree(html)
        warn_unconverted_blocks(root, self.warnings)
        self.render(root)

        self.apply_core_properties()          # 必须在 save 之前写入文档属性
        self.doc.save(str(self.out_path))
        self.scrub_app_properties(self.out_path)
        for td in self.tmp_dirs:
            shutil.rmtree(td, ignore_errors=True)
        if self.verbose:
            print(f"[stats] {self.result.summary()}", file=sys.stderr)
        return self.result


# =============================================================================
# 对外 API（保持旧签名顺序：md_path, out_path, project, bidder, anonymous）
# =============================================================================

def convert(md_path, out_path, project="", bidder="", anonymous=False,
            body_font="仿宋", line_spacing=22, table_style="grid",
            margins=DEFAULT_MARGINS, download_images=False, cover=True,
            verbose=False):
    conv = Converter(md_path, out_path, project=project, bidder=bidder,
                     anonymous=anonymous, body_font=body_font,
                     line_spacing=line_spacing, table_style=table_style,
                     margins=margins, download_images=download_images,
                     cover=cover, verbose=verbose)
    res = conv.build()
    res.out_path = str(out_path)
    return res


def _parse_margin(value):
    if not value:
        return DEFAULT_MARGINS
    parts = [p for p in re.split(r"[,，]", value.strip()) if p.strip()]
    if len(parts) != 4:
        raise argparse.ArgumentTypeError("--margin 需要 4 个数值：上,下,左,右（cm）")
    try:
        nums = tuple(float(p) for p in parts)
    except ValueError:
        raise argparse.ArgumentTypeError("--margin 数值非法")
    if any(n < 0 for n in nums):
        raise argparse.ArgumentTypeError("--margin 不得为负")
    return nums


def main(argv=None):
    ap = argparse.ArgumentParser(description="Markdown → 投标 DOCX（v%s）" % __version__)
    ap.add_argument("md", help="Markdown 文件路径")
    ap.add_argument("--project", default="", help="项目名称（封面用）")
    ap.add_argument("--bidder", default="", help="投标人名称（封面用）")
    ap.add_argument("-o", "--output", default=None, help="输出 DOCX 路径")
    ap.add_argument("--anonymous", action="store_true",
                    help="暗标模式：无封面投标人信息、文档属性清零")
    ap.add_argument("--body-font", default="仿宋", help="正文字体（默认 仿宋）")
    ap.add_argument("--line-spacing", type=int, choices=(22, 28), default=22,
                    help="正文行距磅值（默认 22）")
    ap.add_argument("--table-style", choices=("grid", "three-line"), default="grid",
                    help="表格样式（默认 grid 全框线）")
    ap.add_argument("--margin", default="", type=_parse_margin,
                    help="页边距 上,下,左,右（cm），默认 2.54,2.54,3.17,3.17")
    ap.add_argument("--download-images", action="store_true",
                    help="下载 URL 图片到临时目录（缺省记缺图并 rc=1）")
    ap.add_argument("--no-cover", action="store_true", help="不生成封面与目录")
    ap.add_argument("--verbose", action="store_true", help="打印统计")
    args = ap.parse_args(argv)

    md = Path(args.md)
    if not md.exists():
        print(f"文件不存在: {md}", file=sys.stderr)
        return 2
    out = Path(args.output) if args.output else md.with_suffix(".docx")
    try:
        res = convert(md, out, args.project, args.bidder, args.anonymous,
                      body_font=args.body_font, line_spacing=args.line_spacing,
                      table_style=args.table_style, margins=args.margin,
                      download_images=args.download_images,
                      cover=not args.no_cover, verbose=args.verbose)
    except Exception as exc:                                  # 绝不静默崩溃
        print(f"ERROR: 转换失败：{exc}", file=sys.stderr)
        raise

    for w in res.warnings:
        print(f"[warn] {w}", file=sys.stderr)
    print(f"OK: {out}")
    if args.verbose:
        print(f"     {res.summary()}", file=sys.stderr)
    if not args.anonymous:
        print("提示: 占位符 [资质占位符：xxx] 等需线下确认后手动替换")
    if res.missing_images:
        print(f"ERROR: {len(res.missing_images)} 张图片缺失/不可用，已生成占位（rc=1）：",
              file=sys.stderr)
        for item in res.missing_images:
            print(f"  - {item}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
