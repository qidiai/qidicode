#!/usr/bin/env python3
"""office-canvas: 精修闭环的 craft 自查脚本。

用途:
  在精修（bid-perfect / office-canvas）闭环中做 craft 自查——渲染关键页 PNG
  供 agent 目视核对关键页内容/表格是否落地，并通过 wordcraft-cli 回读结构
  （页数/段落/表格）与文本，替代部分 WPS COM 截图与 python-docx 抽取。

子命令（JSON 输出到 stdout，UTF-8）:
  render <docx> --pages <spec> --out <dir> --scale 2
      --pages 支持 "1,3,5" 显式页 / mid(⌈pages/2⌉) / last / all(上限 20 页)
      页面页码 1-based，每页输出 <dir>/p<N>.png
  text <docx> [--out <file>]
      调 wordcraft-cli text 回读纯文本，并标注修订风险 revision_risk
  inspect <docx>
      调 wordcraft-cli info 回读结构与元数据（pages/paragraphs/tables...）
  --selftest
      python-docx 造 fixture 自检全部子命令

风险引用:
  R19: wordcraft text 可能包含修订模式下已删除的文本（w:delText），因此 text
       子命令返回 revision_risk 标注；revision_risk 非零时不要单凭 text 下结论，
       以 python-docx / canvas.py verify 的结果为准。
  R22: wordcraft 排版与 WPS 存在系统性度量偏差，渲染图仅供结构/内容自查，
       不作为打印级排版比对依据（打印级比对仍走 WPS COM）。

退出码: 0=成功  1=操作失败  2=wordcraft-cli 缺失
"""
import argparse
import glob
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

# ArtCraft 安装根目录；实际 CLI 按 wordcraft-*-windows-x64-portable 通配解析，
# 不硬编码版本号（升级后自动取版本号最大者）。
ARTCRAFT_BIN = r"G:\程序开发\artcraft\bin"
CLI_GLOB = os.path.join(ARTCRAFT_BIN, "wordcraft-*-windows-x64-portable",
                        "wordcraft-cli.exe")
CLI_TIMEOUT = 180          # 单次 CLI 调用超时（秒）
MAX_RENDER_PAGES = 20      # all 模式渲染上限，防爆
MAX_INLINE_TEXT = 20000    # 超过该字符数的 text 只落盘不刷屏
DEFAULT_OUT_DIR = os.path.join(tempfile.gettempdir(), "craft_verify_out")
OUT_DIR_TTL = 24 * 3600    # 默认输出目录中文本/PNG 的保留时长（秒）

RE_VERSION = re.compile(r"wordcraft-([0-9]+(?:\.[0-9]+)*)-windows-x64-portable")

# R19 修订风险计数（独立实现，禁止 import 标书软件代码）
# 字符类含 \n：标签名后紧跟换行（如 "<w:delText\n"）也是合法标记边界，
# 与标书软件 utils/craft_bridge.py 的 R19 三模式保持逐字一致。
RE_REV_DELTEXT = re.compile(r"<w:delText[ >/\n]")
RE_REV_SYM = re.compile(r"<w:sym[ >/\n]")
RE_REV_TXBX = re.compile(r"<w:txbxContent[ >/\n]")


# ---------------------------------------------------------------- CLI 解析

def _version_key(path: str):
    m = RE_VERSION.search(path.replace("\\", "/"))
    if not m:
        return ()
    return tuple(int(x) for x in m.group(1).split("."))


def _resolve_cli():
    """解析 wordcraft-cli 路径：env WORDCRAFT_CLI 优先，否则通配取最高版本。"""
    env = os.environ.get("WORDCRAFT_CLI", "").strip().strip('"')
    if env and os.path.isfile(env):
        return env
    candidates = [p for p in glob.glob(CLI_GLOB) if os.path.isfile(p)]
    if not candidates:
        return None
    candidates.sort(key=_version_key)
    return candidates[-1]


def _run_cli(cli, argv, timeout=CLI_TIMEOUT):
    """执行 CLI，返回 (rc, stdout, stderr, timed_out)。输出按 UTF-8 解码。"""
    try:
        proc = subprocess.run([cli] + argv, capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return -1, "", f"超时({timeout}s)", True
    except OSError as e:
        return -1, "", str(e), False
    out = proc.stdout.decode("utf-8", "replace")
    err = proc.stderr.decode("utf-8", "replace")
    return proc.returncode, out, err, False


def _emit(obj, code=0):
    """JSON 以 UTF-8 字节写 stdout（避免 Windows 控制台代码页破坏中文）。"""
    data = json.dumps(obj, ensure_ascii=False, indent=2).encode("utf-8")
    try:
        sys.stdout.buffer.write(data + b"\n")
        sys.stdout.buffer.flush()
    except AttributeError:  # pragma: no cover - 非常规 stdout
        sys.stdout.write(data.decode("utf-8") + "\n")
    return code


_MISSING_CLI = {"ok": False, "error": "wordcraft-cli 不存在"}


def _require_cli():
    cli = _resolve_cli()
    if cli is None:
        return None, _emit(_MISSING_CLI, 2)
    return cli, None


# ---------------------------------------------------------------- 公共实现

def _check_file(path: str):
    path = os.path.abspath(os.path.expanduser(path))
    if not os.path.isfile(path):
        return None, _emit({"ok": False, "error": f"文件不存在: {path}"}, 1)
    return path, None


def _read_info(cli, docx):
    """调 wordcraft-cli info，返回 (info_dict, error_str)。"""
    rc, out, err, _ = _run_cli(cli, ["info", docx])
    if rc != 0:
        return None, (err or out).strip() or f"info 退出码 {rc}"
    try:
        return json.loads(out), ""
    except json.JSONDecodeError as e:
        return None, f"info 输出非 JSON: {e}: {out[:200]}"


def _read_text(cli, docx):
    """调 wordcraft-cli text，返回 (text, error_str)。"""
    rc, out, err, _ = _run_cli(cli, ["text", docx])
    if rc != 0:
        return "", (err or out).strip() or f"text 退出码 {rc}"
    return out, ""


def _revision_risk(docx):
    """R19: 直读 word/document.xml 统计修订/符号/文本框标记（不 import 标书软件）。"""
    risk = {"deltext_count": 0, "sym_count": 0, "txbx_content_count": 0}
    try:
        with zipfile.ZipFile(docx) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile, OSError) as e:
        risk["error"] = f"无法读取 word/document.xml: {e}"
        risk["revision_present"] = False
        return risk
    risk["deltext_count"] = len(RE_REV_DELTEXT.findall(xml))
    risk["sym_count"] = len(RE_REV_SYM.findall(xml))
    risk["txbx_content_count"] = len(RE_REV_TXBX.findall(xml))
    # P3-3: 与 craft_bridge._revision_markers 对齐——risk 只看 delText/sym；
    # txbx（文本框）仅单列计数，不参与 risk（文本框属已知限制，不等于修订）。
    risk["revision_present"] = (risk["deltext_count"] > 0
                                or risk["sym_count"] > 0)
    return risk


def _cleanup_dir(directory, max_age=OUT_DIR_TTL):
    """清理目录中超过 max_age 秒的文件（每次运行调用；目录不存在不报错）。"""
    try:
        now = time.time()
        with os.scandir(directory) as entries:
            for entry in entries:
                try:
                    if entry.is_file() and now - entry.stat().st_mtime > max_age:
                        os.remove(entry.path)
                except OSError:
                    pass
    except OSError:
        pass


def _text_file_path(docx):
    digest = hashlib.md5(os.path.abspath(docx).encode("utf-8")).hexdigest()[:8]
    stem = os.path.splitext(os.path.basename(docx))[0]
    os.makedirs(DEFAULT_OUT_DIR, exist_ok=True)
    return os.path.join(DEFAULT_OUT_DIR, f"craft_verify_text_{stem}_{digest}.txt")


def _resolve_pages(spec, total_pages):
    """把 --pages 规格解析为 1-based 页码列表（去重保序）。"""
    pages = []
    for token in [t.strip() for t in spec.split(",") if t.strip()]:
        low = token.lower()
        if low == "mid":
            pages.append((total_pages + 1) // 2)
        elif low == "last":
            pages.append(total_pages)
        elif low == "all":
            pages.extend(range(1, min(total_pages, MAX_RENDER_PAGES) + 1))
        else:
            pages.append(int(token))
    seen = set()
    ordered = []
    for n in pages:
        if n not in seen:
            seen.add(n)
            ordered.append(n)
    return ordered


def _do_render(cli, docx, pages_spec, out_dir, scale):
    tokens = [t.strip().lower() for t in pages_spec.split(",") if t.strip()]
    needs_total = any(t in ("mid", "last", "all") for t in tokens)
    total_pages = None
    info, info_err = _read_info(cli, docx)
    if info is not None:
        total_pages = int(info.get("pages") or 0)
    if needs_total:
        if info is None:
            return {"ok": False, "error": info_err}
        if total_pages <= 0:
            return {"ok": False, "error": "info 未返回有效页数"}
    try:
        pages = _resolve_pages(pages_spec, total_pages or 0)
    except ValueError:
        return {"ok": False, "error": f"无法解析 --pages: {pages_spec}"}
    if not pages:
        return {"ok": False, "error": "--pages 未指定任何页"}

    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    start = time.perf_counter()
    pngs, errors = [], []
    # P3-4 显式页码上界预检：info 页数已知时，越界页在调用 CLI 前记入 errors，
    # 不中断其余页（渲染循环只跑 1..total_pages 内的页）。
    render_pages = pages
    if total_pages:
        render_pages = []
        for n in pages:
            if n < 1 or n > total_pages:
                errors.append({
                    "page": n,
                    "error": f"页码越界: {n} 不在 1..{total_pages}（info pages={total_pages}）",
                })
            else:
                render_pages.append(n)
    for n in render_pages:
        png = os.path.join(out_dir, f"p{n}.png")
        rc, out, err, timed_out = _run_cli(
            cli, ["render", docx, png, "--page", str(n), "--scale", str(scale)])
        if rc == 0 and os.path.isfile(png) and os.path.getsize(png) > 0:
            pngs.append(png)
        else:
            detail = (err or out).strip() or ("超时" if timed_out else f"退出码 {rc}")
            errors.append({"page": n, "error": detail[:500]})
    elapsed = int((time.perf_counter() - start) * 1000)

    result = {
        "ok": not errors,
        "pages": pages,
        "pngs": pngs,
        "ms": elapsed,
    }
    if total_pages is not None:
        result["total_pages"] = total_pages
        if any(t == "all" for t in tokens) and total_pages > MAX_RENDER_PAGES:
            result["capped"] = True
    if errors:
        result["errors"] = errors
    return result


def _do_text(cli, docx, out_file=None):
    text, err = _read_text(cli, docx)
    result = {
        "ok": not err,
        "text": text,
        "chars": len(text),
        "revision_risk": _revision_risk(docx),
    }
    if err:
        result["error"] = err
        return result

    text_file = None
    _cleanup_dir(DEFAULT_OUT_DIR)   # P3-5: 每次运行清理默认输出目录中 >24h 文件
    if out_file:
        out_path = os.path.abspath(os.path.expanduser(out_file))
        # --out 指向目录（已存在目录 / 以分隔符结尾）时落该目录，否则按文件路径
        if os.path.isdir(out_path) or out_file.endswith(("/", "\\")):
            os.makedirs(out_path, exist_ok=True)
            text_file = os.path.join(out_path, os.path.basename(_text_file_path(docx)))
        else:
            text_file = out_path
        os.makedirs(os.path.dirname(text_file) or ".", exist_ok=True)
        with open(text_file, "w", encoding="utf-8", newline="") as f:
            f.write(text)
        result["text_file"] = text_file
    if len(text) > MAX_INLINE_TEXT:
        if text_file is None:
            text_file = _text_file_path(docx)   # 专用 tempdir（%TEMP%/craft_verify_out）
            with open(text_file, "w", encoding="utf-8", newline="") as f:
                f.write(text)
            result["text_file"] = text_file
        result["text"] = ""          # 大文本只落盘，不刷屏
        result["text_truncated"] = True
    return result


def _do_inspect(cli, docx):
    info, err = _read_info(cli, docx)
    if info is None:
        return {"ok": False, "error": err}
    text, text_err = _read_text(cli, docx)
    result = {"ok": True, "info": info, "text_chars": len(text) if not text_err else None}
    if text_err:
        result["text_error"] = text_err
        result["text_chars"] = None
    return result


# ---------------------------------------------------------------- 子命令

def cmd_render(args):
    cli, code = _require_cli()
    if cli is None:
        return code
    docx, code = _check_file(args.docx)
    if docx is None:
        return code
    result = _do_render(cli, docx, args.pages, args.out, args.scale)
    return _emit(result, 0 if result.get("ok") else 1)


def cmd_text(args):
    cli, code = _require_cli()
    if cli is None:
        return code
    docx, code = _check_file(args.docx)
    if docx is None:
        return code
    result = _do_text(cli, docx, args.out)
    return _emit(result, 0 if result.get("ok") else 1)


def cmd_inspect(args):
    cli, code = _require_cli()
    if cli is None:
        return code
    docx, code = _check_file(args.docx)
    if docx is None:
        return code
    result = _do_inspect(cli, docx)
    return _emit(result, 0 if result.get("ok") else 1)


# ---------------------------------------------------------------- 自检

def _make_fixture(path):
    """python-docx 造含中文段落 + w:delText 修订删除段的 fixture。"""
    import docx  # 仅 selftest 依赖

    d = docx.Document()
    d.add_heading("Craft Verify 自检", level=1)
    d.add_paragraph("整改测试内容：批注锚定文本已按意见修改。")
    d.add_paragraph("第二段用于验证 render 能生成非空 PNG。")
    d.save(path)

    with zipfile.ZipFile(path) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    xml = entries["word/document.xml"].decode("utf-8")
    inject = (
        '<w:p><w:del w:id="1" w:author="selftest" w:date="2026-10-19T00:00:00Z">'
        "<w:r><w:delText>已删除的修订文本</w:delText></w:r></w:del></w:p>"
    )
    xml = xml.replace("</w:body>", inject + "</w:body>", 1)
    entries["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)


def _make_sym_fixture(path):
    """python-docx 造仅含 w:sym 符号字面量的 fixture（无 delText 修订删除）。"""
    import docx  # 仅 selftest 依赖

    d = docx.Document()
    d.add_paragraph("符号字面量自检：仅有 sym，无修订删除。")
    d.save(path)

    with zipfile.ZipFile(path) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    xml = entries["word/document.xml"].decode("utf-8")
    inject = '<w:p><w:r><w:sym w:font="Wingdings" w:char="F0A7"/></w:r></w:p>'
    xml = xml.replace("</w:body>", inject + "</w:body>", 1)
    entries["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)


def _selftest():
    failures = []

    def check(name, cond, detail=""):
        tag = "PASS" if cond else "FAIL"
        print(f"{tag} {name}" + (f" ({detail})" if detail else ""))
        if not cond:
            failures.append(name)
        return cond

    cli = _resolve_cli()
    if not check("wordcraft-cli 解析", cli is not None, cli or "不存在"):
        print("SELFTEST FAILED: wordcraft-cli 不存在")
        return 1

    # P3-1/P3-2 正则边界用例（纯内存，无需 CLI）
    deltext_nl = "<w:delText\n"
    txbx_bad = "<w:txbxContentX"
    check("R19 '<w:delText\\n' 计数=1",
          len(RE_REV_DELTEXT.findall(deltext_nl)) == 1,
          f"count={len(RE_REV_DELTEXT.findall(deltext_nl))}")
    check("R19 '<w:txbxContentX' 计数=0",
          len(RE_REV_TXBX.findall(txbx_bad)) == 0,
          f"count={len(RE_REV_TXBX.findall(txbx_bad))}")

    workdir = tempfile.mkdtemp(prefix="craft_verify_selftest_")
    try:
        fixture = os.path.join(workdir, "fixture.docx")
        _make_fixture(fixture)
        check("fixture 生成", os.path.isfile(fixture), fixture)

        render = _do_render(cli, fixture, "1", os.path.join(workdir, "png"), 2)
        pngs = render.get("pngs") or []
        png_size = os.path.getsize(pngs[0]) if pngs else 0
        check("render p1 ok", bool(render.get("ok")) and len(pngs) == 1,
              f"ok={render.get('ok')} pngs={len(pngs)}")
        check("render p1 PNG>1KB", png_size > 1024, f"{png_size} bytes")

        text_result = _do_text(cli, fixture)
        check("text 回读", bool(text_result.get("ok"))
              and "整改测试内容" in text_result.get("text", ""),
              f"chars={text_result.get('chars')}")

        risk = text_result.get("revision_risk") or {}
        check("revision_risk.deltext_count>0",
              int(risk.get("deltext_count") or 0) > 0,
              f"deltext_count={risk.get('deltext_count')}")

        # P3-3 仅 sym 的 fixture -> revision_present=True
        sym_fixture = os.path.join(workdir, "sym_fixture.docx")
        _make_sym_fixture(sym_fixture)
        sym_risk = _revision_risk(sym_fixture)
        check("P3-3 仅 sym -> revision_present=True",
              int(sym_risk.get("sym_count") or 0) > 0
              and int(sym_risk.get("deltext_count") or 0) == 0
              and sym_risk.get("revision_present") is True,
              f"sym_count={sym_risk.get('sym_count')} "
              f"deltext_count={sym_risk.get('deltext_count')} "
              f"revision_present={sym_risk.get('revision_present')}")

        # P3-4 越界页预检：page 99 记 errors，page 1 照常渲染
        oob = _do_render(cli, fixture, "1,99",
                         os.path.join(workdir, "png_oob"), 2)
        oob_errors = [e.get("page") for e in (oob.get("errors") or [])]
        check("P3-4 越界页预检(99 记错/1 照渲染)",
              oob.get("ok") is False and 99 in oob_errors
              and len(oob.get("pngs") or []) == 1,
              f"error_pages={oob_errors} pngs={len(oob.get('pngs') or [])}")

        # P3-5 每次运行清理默认输出目录中 >24h 的文件
        os.makedirs(DEFAULT_OUT_DIR, exist_ok=True)
        stale = os.path.join(DEFAULT_OUT_DIR, "stale_selftest.txt")
        with open(stale, "w", encoding="utf-8") as f:
            f.write("stale")
        old_ts = time.time() - 25 * 3600
        os.utime(stale, (old_ts, old_ts))
        _do_text(cli, fixture)
        check("P3-5 清理 >24h 落盘文件", not os.path.exists(stale), stale)

        inspect_result = _do_inspect(cli, fixture)
        check("inspect ok", bool(inspect_result.get("ok"))
              and isinstance(inspect_result.get("info"), dict),
              f"pages={(inspect_result.get('info') or {}).get('pages')}")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)

    if failures:
        print("SELFTEST FAILED: " + ", ".join(failures))
        return 1
    print("SELFTEST OK")
    return 0


# ---------------------------------------------------------------- 入口

def main(argv=None):
    parser = argparse.ArgumentParser(
        description="office-canvas 精修闭环 craft 自查（wordcraft-cli 封装）")
    parser.add_argument("--selftest", action="store_true",
                        help="用 python-docx 造 fixture 自检全部子命令")
    sub = parser.add_subparsers(dest="cmd")

    p_render = sub.add_parser("render", help="渲染关键页 PNG")
    p_render.add_argument("docx")
    p_render.add_argument("--pages", default="1,mid,last",
                          help="1,3,5 / mid / last / all（可混用，上限 20 页）")
    p_render.add_argument("--out", default=DEFAULT_OUT_DIR,
                          help="PNG 输出目录（默认 %%TEMP%%/craft_verify_out = %s）"
                               % DEFAULT_OUT_DIR)
    p_render.add_argument("--scale", type=float, default=2.0)
    p_render.set_defaults(func=cmd_render)

    p_text = sub.add_parser("text", help="文本回读 + 修订风险标注")
    p_text.add_argument("docx")
    p_text.add_argument("--out", default=DEFAULT_OUT_DIR,
                        help="文本落盘路径或目录（默认 %%TEMP%%/craft_verify_out = %s；"
                             "传目录时落 <dir>/craft_verify_text_*.txt）" % DEFAULT_OUT_DIR)
    p_text.set_defaults(func=cmd_text)

    p_inspect = sub.add_parser("inspect", help="结构回读(页数/段落/表格)")
    p_inspect.add_argument("docx")
    p_inspect.set_defaults(func=cmd_inspect)

    args = parser.parse_args(argv)
    if args.selftest:
        return _selftest()
    if not getattr(args, "func", None):
        parser.print_help()
        return 1
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
